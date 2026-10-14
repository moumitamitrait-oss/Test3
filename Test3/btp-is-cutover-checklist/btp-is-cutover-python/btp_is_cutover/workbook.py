"""The cutover workbook: List of Dependent Artifacts, CutOver_<ENV>_<Interface ID>, Config Params."""
from __future__ import annotations

import math
import os
import re
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import CutoverData
from .steps import build_steps

GREEN, LGREEN, LBLUE, HDR, SUB, SHEET1 = "70AD47", "E2EFDA", "D9E1F2", "00B0F0", "8EA9DB", "9BC2E6"
_thin = Side(style="thin", color="000000")
BORDER = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)


def _font(**kw) -> Font:
    return Font(name="Calibri", size=kw.pop("size", 11), **kw)


def _fill(rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=rgb)


def _lines(text: Any, width: int) -> int:
    return max(1, math.ceil(len(str(text or "")) / width))


def _widths(ws, widths: List[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _safe(text: str) -> str:
    return re.sub(r"[\\/?*\[\]:]", "_", str(text))


def sheet_dependent(ws, data: CutoverData, common: List[Dict[str, str]]) -> None:
    for c, h in enumerate(["S.No", "Package", "Artifact", "System"], 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill = _font(bold=True, color="000000"), _fill(SHEET1)
    rows = [(d["pkg"], d["artifact"], d["system"]) for d in common]
    rows.append((data.iflow["Package"], data.iflow["Id"], data.sender_system or ""))
    for i, (pkg, art, system) in enumerate(rows, 1):
        for c, v in enumerate([i, pkg, art, system], 1):
            ws.cell(i + 1, c, v).font = _font()
    _widths(ws, [4.6, 42.7, 44.6, 29.7])


def sheet_cutover(ws, rows: List[Dict[str, Any]]) -> None:
    heads = ["Step", "Reference Artifact", "Integration Package", "Actions", "Task execution time", "Notes",
             "Assignee Team", "Point Of Contact", "Status", "Completion Date"]
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c, h in enumerate(heads, 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill, cell.alignment, cell.border = _font(bold=True, color="000000"), _fill(HDR), center, BORDER
    ws.row_dimensions[1].height = 14.5
    for i, r in enumerate(rows, 2):
        if r["kind"] == "banner":
            for c in range(1, 11):
                cell = ws.cell(i, c, r["text"] if c == 4 else "")
                cell.font, cell.fill, cell.alignment = _font(bold=True, size=14, color="000000"), _fill(GREEN), center
            ws.row_dimensions[i].height = 23.5
            continue
        values = [r["n"], r["ref"], r["pkg"], r["action"], r.get("time", ""), r.get("notes", ""), r["team"], r["poc"], "Not Started", ""]
        for c, v in enumerate(values, 1):
            cell = ws.cell(i, c, v)
            cell.font, cell.fill, cell.alignment, cell.border = _font(), _fill(LBLUE if r.get("smoke") else LGREEN), center, BORDER
        n = max(_lines(r["action"], 40), _lines(r["ref"], 30), _lines(r["pkg"], 42), _lines(r.get("notes"), 36))
        ws.row_dimensions[i].height = max(29, n * 15)
    _widths(ws, [20.5, 36.3, 51, 54.7, 22.5, 43, 27.2, 29.3, 34.7, 18.2])


def sheet_config(ws, data: CutoverData) -> None:
    """Parameter Name | DEV | QA | PROD. Each parameter once, in the order Sender -> Receiver -> More."""
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for c, h in enumerate(["Parameter Name", "DEV", "QA", "PROD"], 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill, cell.alignment, cell.border = _font(bold=True, color="000000"), _fill(SUB), left, BORDER
    ws.row_dimensions[1].height = 19
    seen = set()
    items: List[Tuple[str, str]] = []
    for g in sorted(data.param_groups, key=lambda g: {"Sender": 0, "Receiver": 1, "More": 2}[g["category"]]):
        for key, value in g["items"]:
            if key not in seen:
                seen.add(key)
                items.append((key, value))
    for i, (key, value) in enumerate(items, 2):
        for c, v in enumerate([key, value, "", ""], 1):
            cell = ws.cell(i, c, v[:32000] if isinstance(v, str) else v)
            cell.font, cell.alignment, cell.border = _font(), left, BORDER
        n = max(_lines(value, 52), _lines(key, 42))
        if n > 1:
            ws.row_dimensions[i].height = n * 15
    _widths(ws, [46, 56, 40, 40])


def build_workbook(data: CutoverData, options: Dict[str, Any]) -> Tuple[Workbook, str]:
    """options: env, interface_id, poc, team_is, team_func, team_sec, common (list of {pkg, artifact, system})."""
    env = options.get("env") or "QA"
    label = (options.get("interface_id") or data.iflow["Id"]).strip()
    common = [c for c in (options.get("common") or []) if c.get("artifact")]
    rows = build_steps(data, options)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "List of Dependent Artifacts"
    sheet_dependent(ws1, data, common)
    ws2 = wb.create_sheet(_safe(f"CutOver_{env}_{label}")[:31])
    sheet_cutover(ws2, rows)
    ws3 = wb.create_sheet("Config Params")
    sheet_config(ws3, data)

    system = re.sub(r"[^\w.\-]", "_", data.sender_system or "")
    filename = f"CutOver_{_safe(env)}_{_safe(label).replace('.', '_')}{'_' + system if system else ''}.xlsx"
    return wb, filename


def save_workbook(wb: Workbook, folder: str, filename: str) -> str:
    os.makedirs(folder, exist_ok=True)
    path, n = os.path.join(folder, filename), 2
    while os.path.exists(path):  # never overwrite an earlier file
        path = os.path.join(folder, re.sub(r"\.xlsx$", "", filename) + f"_{n}.xlsx")
        n += 1
    wb.save(path)
    return path

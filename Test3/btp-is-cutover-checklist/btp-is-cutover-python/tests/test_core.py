import os

import pytest
from openpyxl import load_workbook

from btp_is_cutover.api import ApiError, CpiClient, get_token, lit, normalize_base
from btp_is_cutover.cutover import cutover_data_for
from btp_is_cutover.design import parse_design, parse_manifest
from btp_is_cutover.detect import detect_dependencies, hint_for, with_defaults
from btp_is_cutover.models import Dep
from btp_is_cutover.service import analyze, generate, validate_details
from btp_is_cutover.steps import build_steps
from btp_is_cutover.tenant import SelectedFlow, Tenant
from btp_is_cutover.workbook import build_workbook
from tests import fake_tenant as ft

BASE = "https://demo.it-cpi001.cfapps.eu10.hana.ondemand.com"


@pytest.fixture
def session():
    return ft.FakeSession()


@pytest.fixture
def tenant(session):
    return Tenant.connect(BASE, "https://sub.authentication.eu10.hana.ondemand.com", "id", "secret", session=session)


def flow(tenant, iflow_id):
    f = tenant.resolve_iflow_by_id(iflow_id)
    f.interface_id = "RICEFW-1"
    return f.as_flow()


# ---- api -------------------------------------------------------------------------------------------------------
def test_helpers():
    assert lit("it's") == "'it''s'" or lit("it's") == "'it%27%27s'" or "it" in lit("it's")
    assert normalize_base("https://x.com/") == "https://x.com/api/v1"
    assert normalize_base("https://x.com/api/v1") == "https://x.com/api/v1"


def test_token_errors(session):
    assert get_token("https://t", "id", "secret", session=session) == "TOKEN"
    with pytest.raises(ApiError) as e:
        get_token("https://t", "id", "wrong", session=session)
    assert e.value.status == 401 and "client ID" in str(e.value)


def test_501_is_not_retried_and_error_names_the_call(session):
    session.fail["/Configurations"] = 501
    client = CpiClient(BASE, "T", session=session, sleep=lambda s: None)
    with pytest.raises(ApiError) as e:
        client.get_configurations({"Id": "IF_OSS", "Version": "1.0.3"})
    assert e.value.status == 501 and "Configurations" in str(e.value)
    assert sum(1 for c in session.calls if "Configurations" in c) == 1


def test_503_is_retried(session):
    n = {"c": 0}
    real = session.get

    def flaky(url, headers=None, timeout=None):
        if "IntegrationPackages" in url and n["c"] < 2:
            n["c"] += 1
            return ft.Resp(503, {})
        return real(url, headers, timeout)

    session.get = flaky
    client = CpiClient(BASE, "T", session=session, sleep=lambda s: None)
    assert len(client.get_all_paged("/IntegrationPackages")) == 4 and n["c"] == 2


def test_version_requested_and_missing(tenant):
    f = flow(tenant, "IF_OSS")
    cfg, v = tenant.client.get_configurations({**f, "reqVersion": "1.0.3"})
    assert v == "1.0.3" and len(cfg) == 6
    with pytest.raises(ApiError) as e:
        tenant.client.get_configurations({**f, "reqVersion": "9.9.9"})
    assert "not found" in str(e.value)


# ---- tenant ----------------------------------------------------------------------------------------------------
def test_resolve_by_id_direct_and_fallback(tenant, session):
    f = tenant.resolve_iflow_by_id("IF_OSS")
    assert (f.pkg_name, f.id) == ("P66_HVC_App", "IF_OSS")
    session.direct_status = 404
    session.calls.clear()
    g = tenant.resolve_iflow_by_id("if_amqp")            # falls back to looking through the packages
    assert g.id == "IF_AMQP" and any("IntegrationPackages" in c for c in session.calls)
    with pytest.raises(ApiError):
        tenant.resolve_iflow_by_id("NOT_THERE")


def test_forbidden_direct_read_is_reported(tenant, session):
    session.fail["IntegrationDesigntimeArtifacts(Id='IF_OSS'"] = 403
    with pytest.raises(ApiError) as e:
        tenant.resolve_iflow_by_id("IF_OSS")
    assert "403" in str(e.value) and "roles" in str(e.value)


# ---- design ----------------------------------------------------------------------------------------------------
def test_manifest_refs():
    assert parse_manifest(ft.MM_WITH["META-INF/MANIFEST.MF"]) == ["Global_Lib_Artifact"]
    assert parse_manifest("") == []


def test_design_channels_credentials_timer():
    d = parse_design(ft.IF_OSS)
    types = [(c["direction"], c["type"], c["participant"]) for c in d.channels]
    assert ("Sender", "SFTP", "S_OSS") in types and ("Receiver", "HTTP", "R_SAPS4HANA") in types
    assert d.usage["S4_Host"] == [("Receiver", "R_SAPS4HANA")]
    assert not d.timer and parse_design(ft.IF_TIMER).timer
    assert any(c["address"] == "{{PD_Url}}" for c in d.channels)


# ---- cutover data ----------------------------------------------------------------------------------------------
def test_systems_and_credentials(tenant):
    data = cutover_data_for(tenant, flow(tenant, "IF_OSS"))
    assert data.sender_system == "OSS" and data.has_sender and data.sender_type == "SFTP"
    assert data.receiver_system == "SAPS4HANA" and data.receiver_credential == "SAP_S4_110_PROXY_Cred"
    assert [c["alias"] for c in data.credentials] == ["SAP_S4_110_PROXY_Cred"]      # no "false" / "Basic" junk
    assert "AlertNotification" not in data.receiver_system


def test_pd_only_and_timer(tenant):
    pd = cutover_data_for(tenant, flow(tenant, "IF_PDONLY"))
    assert not pd.has_sender and pd.sender_system == ""
    tm = cutover_data_for(tenant, flow(tenant, "IF_TIMER"))
    assert tm.timer and tm.sender_system == "Timer" and not tm.has_sender


def test_design_unreadable_gives_warning_not_failure(tenant, session):
    session.fail["/$value"] = 403
    data = cutover_data_for(tenant, flow(tenant, "IF_OSS"))
    assert data.warnings and "could not be read" in data.warnings[0]
    assert [g["category"] for g in data.param_groups] == ["More"]


# ---- dependencies ----------------------------------------------------------------------------------------------
def deps_of(tenant, iflow_id="IF_OSS"):
    data = cutover_data_for(tenant, flow(tenant, iflow_id))
    return data, detect_dependencies(tenant, data)


def test_hints():
    assert hint_for("scriptBundleId") == "Script Collection" and hint_for("mappingname") == "Message Mapping"
    assert hint_for("valueMappingId") == "Value Mapping" and hint_for("sourceMessageType") == "Message Type"
    assert hint_for("something") is None


def test_dependencies_found(tenant):
    _, res = deps_of(tenant)
    got = {(d.artifact, d.system) for d in res.deps}
    assert ("SC_Global_ReusableScripts", "Script Collection") in got           # type from the property name
    assert ("MM_WithGlobals", "Message Mapping") in got and ("MM_Plain", "Message Mapping") in got
    assert ("IF_Global_ErrorLogging_AlertEmail_ServiceNow_AIF", "Integration Flow") in got   # ProcessDirect target
    # global resources used inside the message mapping
    assert ("VM_Global_IFlow_to_EMailRecipients", "Value Mapping") in got and ("Global_Lib_Artifact", "Script Collection") in got
    assert ("MT_Order_Source", "Message Type") in got and ("MT_Common_Header", "Message Type") in got
    assert any(d.artifact == "MT_Order_Target" and d.pkg == "" for d in res.deps)         # clearly a message type, package unknown
    assert not any(d.artifact == "static" for d in res.deps)                            # mappingReference = static is a setting
    assert "SC_Not_In_Tenant" in res.skipped
    joined = " ".join(res.notes)
    assert "MM_Plain: scanned, no references" in joined


def test_dependencies_without_direct_reads(tenant, session):
    session.direct_status = 501                                  # tenant that rejects direct artifact reads
    _, res = deps_of(tenant)
    got = {(d.artifact, d.system) for d in res.deps}
    assert ("SC_Global_ReusableScripts", "Script Collection") in got and ("MM_WithGlobals", "Message Mapping") in got


def test_unopenable_mapping_is_reported(tenant):
    _, res = deps_of(tenant)
    assert any("MM_Locked" in n and "could not be opened" in n for n in res.notes)


def test_defaults_added_once():
    deps = with_defaults([])
    assert deps[0].artifact == "VM_Global_IFlow_to_EMailRecipients" and deps[0].include
    again = with_defaults(deps)
    assert len(again) == 1


# ---- steps -----------------------------------------------------------------------------------------------------
def steps_text(tenant, iflow_id, common=()):
    data = cutover_data_for(tenant, flow(tenant, iflow_id))
    rows = build_steps(data, {"env": "QA", "poc": "J. Doe", "team_is": "IS", "team_func": "Func", "team_sec": "Sec", "common": list(common)})
    return data, [r for r in rows if r["kind"] == "step"]


def test_sftp_sender_steps(tenant):
    _, steps = steps_text(tenant, "IF_OSS")
    text = [s["action"] for s in steps]
    sched = [s for s in steps if "scheduler timings" in s["action"]]
    assert len(sched) == 1 and sched[0]["team"] == "Func" and "SFTP" in sched[0]["notes"]
    assert "Share the Sender SFTP directory and file details to OSS Team" in text
    assert "OSS Team will change the directory and file details from Dev to QA" in text
    assert "Ensure the service user used by credential SAP_S4_110_PROXY_Cred has access to post the data in SAPS4HANA QA" in text
    assert not any("AlertNotification" in t or "ProcessDirect" in t or "(false)" in t for t in text)


def test_processdirect_sender_gets_no_sender_steps(tenant):
    _, steps = steps_text(tenant, "IF_PDONLY")
    assert not any("Share the Sender" in s["action"] or "Team will change" in s["action"] for s in steps)
    assert not any("scheduler" in s["action"] for s in steps)


def test_timer_gets_scheduler_step_only(tenant):
    _, steps = steps_text(tenant, "IF_TIMER")
    assert any("scheduler timings" in s["action"] and s["notes"] == "Timer start event" for s in steps)
    assert not any("Share the Sender" in s["action"] for s in steps)
    assert not any("service user" in s["action"] for s in steps)           # no receiver credential


def test_amqp_queue_wording_and_email_steps(tenant):
    common = [{"pkg": "P66_Global_GenericHandlers", "artifact": "VM_Global_IFlow_to_EMailRecipients", "system": "Value Mapping"}]
    _, steps = steps_text(tenant, "IF_AMQP", common)
    text = [s["action"] for s in steps]
    assert "Share the Sender AMQP queue details to ACTIAN Team" in text
    assert "Receive List of Email recipients from Business" in text
    assert any(t.startswith("Configure the respective mail Id") for t in text)
    assert text[-1] == "Testing the Interface End to End"


def test_validation_texts_per_type(tenant):
    common = [{"pkg": "P", "artifact": "SC_X", "system": "Script Collection"}, {"pkg": "P", "artifact": "MT_A", "system": "Message Type"},
              {"pkg": "P", "artifact": "DT_A", "system": "Data Type"}, {"pkg": "P", "artifact": "MM_A", "system": "Message Mapping"},
              {"pkg": "P", "artifact": "IF_A", "system": "Integration Flow"}]
    _, steps = steps_text(tenant, "IF_AMQP", common)
    text = [s["action"] for s in steps]
    for expected in ["Validate the SC_X deployed or not", "Validate the Message Type MT_A deployed or not", "Validate the Data Type DT_A deployed or not",
                     "Validate the Message Mapping MM_A deployed or not", "Validate IF_A flow Deployed or not"]:
        assert expected in text


# ---- workbook / service ----------------------------------------------------------------------------------------
def test_workbook_layout(tenant, tmp_path):
    data = cutover_data_for(tenant, flow(tenant, "IF_OSS"))
    common = [{"pkg": "P66_Global_GenericHandlers", "artifact": "VM_Global_IFlow_to_EMailRecipients", "system": "Value Mapping"}]
    wb, name = build_workbook(data, {"env": "QA", "interface_id": "PTP-I-514.1", "poc": "J", "common": common})
    assert name == "CutOver_QA_PTP-I-514_1_OSS.xlsx"
    assert wb.sheetnames == ["List of Dependent Artifacts", "CutOver_QA_PTP-I-514.1", "Config Params"]
    cfg = wb["Config Params"]
    assert [c.value for c in cfg[1]] == ["Parameter Name", "DEV", "QA", "PROD"] and not cfg.merged_cells.ranges
    keys = [r[0].value for r in cfg.iter_rows(min_row=2)]
    assert len(keys) == len(set(keys)) == 6
    assert dict((r[0].value, r[1].value) for r in cfg.iter_rows(min_row=2))["Unused_Flag"] == "(empty)"
    assert cfg["C2"].value in (None, "") and cfg["D2"].value in (None, "")
    dep = wb["List of Dependent Artifacts"]
    assert [c.value for c in dep[3]][1:] == ["P66_HVC_App", "IF_OSS", "OSS"]      # the iFlow itself, system = sender system
    assert wb.worksheets[1]["A1"].fill.fgColor.rgb.endswith("00B0F0")


def test_service_analyze_and_generate(tenant, tmp_path):
    sel = [tenant.resolve_iflow_by_id("IF_OSS"), tenant.resolve_iflow_by_id("IF_AMQP")]
    sel[0].interface_id, sel[1].interface_id = "RICEFW-1", "RICEFW-2"
    assert validate_details(sel, "") == ["the Point of contact"]
    sel[1].interface_id = ""
    assert "an Interface ID for every iFlow" in validate_details(sel, "J")
    sel[1].interface_id = "RICEFW-2"
    assert validate_details(sel, "J") == []
    res = analyze(tenant, sel)
    assert any(d.artifact == "VM_Global_IFlow_to_EMailRecipients" and d.include for d in res[sel[0].key].result.deps)   # default value mapping
    paths, notes = generate(res, sel, {"env": "QA", "poc": "J"}, str(tmp_path))
    assert [os.path.basename(p) for p in paths] == ["CutOver_QA_RICEFW-1_OSS.xlsx", "CutOver_QA_RICEFW-2_ACTIAN.xlsx"]
    wb = load_workbook(paths[0])
    actions = [r[3].value for r in wb.worksheets[1].iter_rows(min_row=2) if isinstance(r[0].value, int)]
    assert "Receive List of Email recipients from Business" in actions            # because of the default value mapping


def test_unticked_dependencies_are_left_out(tenant, tmp_path):
    sel = [tenant.resolve_iflow_by_id("IF_AMQP")]
    sel[0].interface_id = "RICEFW-9"
    res = analyze(tenant, sel)
    for d in res[sel[0].key].result.deps:
        d.include = False
    paths, _ = generate(res, sel, {"env": "QA", "poc": "J"}, str(tmp_path))
    wb = load_workbook(paths[0])
    assert wb.worksheets[0].max_row == 2                       # only the iFlow itself


def test_failed_analysis_is_skipped_with_a_note(tenant, session, tmp_path):
    sel = [tenant.resolve_iflow_by_id("IF_AMQP")]
    sel[0].interface_id = "R"
    session.fail["/Configurations"] = 403
    res = analyze(tenant, sel)
    paths, notes = generate(res, sel, {"env": "QA", "poc": "J"}, str(tmp_path))
    assert paths == [] and "skipped" in notes[0]

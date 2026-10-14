import json
import os

import pytest
from openpyxl import load_workbook

from btp_is_cutover.cli import main
from tests import fake_tenant as ft

ARGS = ["--base-url", "https://demo.it-cpi001.cfapps.eu10.hana.ondemand.com", "--token-url", "https://demo.authentication.eu10.hana.ondemand.com", "--client-id", "id"]


@pytest.fixture(autouse=True)
def secret(monkeypatch):
    monkeypatch.setenv("BTP_IS_CLIENT_SECRET", "secret")


def test_packages_and_iflows(capsys):
    assert main(["packages"] + ARGS, session=ft.FakeSession()) == 0
    out = capsys.readouterr().out
    assert "P_GLOBAL\tP66_Global_GenericHandlers" in out
    assert main(["iflows", "--package", "p66_hvc_app"] + ARGS, session=ft.FakeSession()) == 0
    assert "IF_OSS\tIF_OSS_to_SAPS4HANA\tv1.0.3" in capsys.readouterr().out
    assert main(["iflows", "--package", "nothing"] + ARGS, session=ft.FakeSession()) == 1


def write_cfg(tmp_path, **over):
    cfg = {"env": "QA", "poc": "Jane Doe", "team_is": "BTP-IS Team", "output_dir": str(tmp_path / "out"),
           "iflows": [{"id": "IF_OSS", "interface_id": "RICEFW-1"}, {"id": "IF_AMQP", "interface_id": "RICEFW-2"}]}
    cfg.update(over)
    path = tmp_path / "cfg.json"
    path.write_text(json.dumps(cfg))
    return str(path)


def test_create_from_config(tmp_path, capsys):
    cfg = write_cfg(tmp_path, extra_dependencies=[{"package": "P_X", "artifact": "VM_Extra", "type": "Value Mapping"}],
                    exclude_dependencies=["MM_Plain"])
    assert main(["create", "--config", cfg] + ARGS, session=ft.FakeSession()) == 0
    out = capsys.readouterr().out
    assert "Created" in out and "RICEFW-1" in out
    files = sorted(os.listdir(tmp_path / "out"))
    assert files == ["CutOver_QA_RICEFW-1_OSS.xlsx", "CutOver_QA_RICEFW-2_ACTIAN.xlsx"]
    wb = load_workbook(tmp_path / "out" / files[0])
    artifacts = [r[2].value for r in wb.worksheets[0].iter_rows(min_row=2)]
    assert "VM_Extra" in artifacts and "MM_Plain" not in artifacts and "SC_Global_ReusableScripts" in artifacts


def test_mandatory_fields_are_enforced(tmp_path, capsys):
    assert main(["create", "--config", write_cfg(tmp_path, poc="")] + ARGS, session=ft.FakeSession()) == 2
    assert "Point of contact" in capsys.readouterr().err
    cfg = write_cfg(tmp_path, iflows=[{"id": "IF_OSS"}])
    assert main(["create", "--config", cfg] + ARGS, session=ft.FakeSession()) == 2
    assert "Interface ID" in capsys.readouterr().err


def test_connection_errors_are_readable(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("BTP_IS_CLIENT_SECRET", "wrong")
    assert main(["packages"] + ARGS, session=ft.FakeSession()) == 1
    assert "Token request failed: HTTP 401" in capsys.readouterr().err
    monkeypatch.delenv("BTP_IS_CLIENT_SECRET")
    assert main(["packages"] + ARGS, session=ft.FakeSession()) == 1
    assert "required" in capsys.readouterr().err
